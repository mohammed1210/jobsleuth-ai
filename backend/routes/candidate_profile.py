"""Candidate CV profile routes.

Raw CV files are processed transiently and are not stored. Only a structured
candidate profile is persisted for vacancy matching.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import APIRouter, File, Header, HTTPException, UploadFile
from pydantic import BaseModel, Field

from lib.settings import settings
from lib.supabase import get_supabase_client
from routes.saved_jobs import verify_supabase_user

router = APIRouter(prefix="/candidate-profile", tags=["candidate_profile"])

MAX_CV_BYTES = 5 * 1024 * 1024
ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt"}


class CandidateExperience(BaseModel):
    role: str = ""
    organisation: str = ""
    dates: str = ""
    highlights: list[str] = Field(default_factory=list)
    skills: list[str] = Field(default_factory=list)


class CandidateQualification(BaseModel):
    name: str
    institution: str = ""
    date: str = ""


class CandidateProfileData(BaseModel):
    summary: str = ""
    skills: list[str] = Field(default_factory=list)
    experience: list[CandidateExperience] = Field(default_factory=list)
    qualifications: list[CandidateQualification] = Field(default_factory=list)


class CandidateProfileResponse(CandidateProfileData):
    user_id: str
    source_filename: str = ""
    extraction_provider: str = "fallback"
    updated_at: datetime | str | None = None


def _db() -> Any:
    return get_supabase_client()


def _dedupe(values: list[str], *, limit: int = 40) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = re.sub(r"\s+", " ", str(value)).strip(" \t\n•·,-")
        key = cleaned.casefold()
        if cleaned and key not in seen:
            seen.add(key)
            output.append(cleaned[:180])
        if len(output) >= limit:
            break
    return output


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    """Extract readable DOCX text, including text boxes and shape content.

    python-docx exposes normal paragraphs and tables but not all DrawingML/text-box
    content. Many CV templates use those shapes heavily, so also walk WordprocessingML
    paragraph nodes from the DOCX archive and prefer that richer representation.
    """
    from docx import Document

    document = Document(io.BytesIO(data))
    standard_parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                standard_parts.extend(
                    paragraph.text
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )

    xml_parts: list[str] = []
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml_names = [
            name
            for name in archive.namelist()
            if name == "word/document.xml"
            or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
        ]
        for name in xml_names:
            root = ET.fromstring(archive.read(name))
            for paragraph in root.iter(f"{namespace}p"):
                pieces = [
                    node.text or ""
                    for node in paragraph.iter(f"{namespace}t")
                    if (node.text or "").strip()
                ]
                line = "".join(pieces).strip()
                if line:
                    xml_parts.append(line)

    standard_text = "\n".join(standard_parts)
    xml_text = "\n".join(xml_parts)
    return xml_text if len(xml_text) > len(standard_text) else standard_text


def _extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            text = _extract_pdf(data)
        elif suffix == ".docx":
            text = _extract_docx(data)
        elif suffix == ".txt":
            text = data.decode("utf-8", errors="replace")
        else:
            raise HTTPException(status_code=415, detail="Upload a PDF, DOCX or TXT CV.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Could not read this CV file. Try a standard PDF, DOCX or TXT file.") from exc

    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise HTTPException(status_code=422, detail="Could not extract enough readable text from this CV.")
    return text[:50000]


_SECTION_RE = re.compile(
    r"(?im)^\s*(skills|skills profile|key skills|technical skills|core skills|key competencies|core competencies|skills & competencies|"
    r"experience|employment|employment history|work history|career history|work experience|professional experience|career experience|"
    r"qualifications|education|education/training|education & training|education & qualifications|qualifications & education|certifications|professional qualifications|"
    r"profile|personal profile|professional profile|objective|career objective|summary|professional summary|about me)\s*:?\s*$"
)


def _sections(text: str) -> dict[str, str]:
    matches = list(_SECTION_RE.finditer(text))
    if not matches:
        return {}
    output: dict[str, str] = {}
    for index, match in enumerate(matches):
        heading = match.group(1).lower()
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        output[heading] = text[start:end].strip()
    return output


_MONTH = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
_DATE_RANGE_RE = re.compile(
    rf"(?i)^\s*{_MONTH}\s+\d{{4}}\s*[–—-]\s*(?:present|current|{_MONTH}\s+\d{{4}})\s*$"
)
_STOP_HEADINGS = {
    "education",
    "education/training",
    "education & training",
    "education & qualifications",
    "qualifications",
    "qualifications & education",
    "certifications",
    "professional qualifications",
    "interests",
    "references",
}


def _clean_cv_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip(" \t•·-")
        if not line:
            continue
        lowered = line.casefold()
        if re.match(r"^(?:d\.?\s*o\.?\s*b\.?|date of birth)\s*:", lowered):
            continue
        if "@" in line and re.search(r"\b(?:tel|phone|mobile|email|e-mail)\b", lowered):
            continue
        lines.append(line)
    return lines


def _is_job_date_line(value: str) -> bool:
    return bool(_DATE_RANGE_RE.fullmatch(value.strip()))


def _split_role_organisation(value: str) -> tuple[str, str]:
    if ":" not in value:
        return value.strip(), ""
    role, organisation = value.split(":", 1)
    return role.strip(), organisation.strip()


def _experience_from_role_date_pairs(lines: list[str], skills: list[str]) -> list[CandidateExperience]:
    experience: list[CandidateExperience] = []
    index = 0
    while index + 1 < len(lines):
        role_line = lines[index]
        lowered = role_line.casefold().rstrip(":")
        if lowered in _STOP_HEADINGS:
            if lowered.startswith("education") or "qualification" in lowered:
                break
            index += 1
            continue

        date_line = lines[index + 1]
        if not _is_job_date_line(date_line):
            index += 1
            continue

        role, organisation = _split_role_organisation(role_line)
        if len(role) < 3 or len(role) > 120:
            index += 1
            continue

        highlights: list[str] = []
        cursor = index + 2
        while cursor < len(lines):
            current = lines[cursor]
            current_lower = current.casefold().rstrip(":")
            if current_lower in _STOP_HEADINGS:
                break
            if cursor + 1 < len(lines) and _is_job_date_line(lines[cursor + 1]):
                break
            highlights.append(current)
            cursor += 1

        experience.append(
            CandidateExperience(
                role=role,
                organisation=organisation,
                dates=date_line,
                highlights=_dedupe(highlights, limit=20),
                skills=skills[:12],
            )
        )
        index = max(cursor, index + 2)
    return experience


def _fallback_profile(text: str) -> CandidateProfileData:
    sections = _sections(text)
    lines = _clean_cv_lines(text)

    skills_text = next(
        (
            sections[key]
            for key in (
                "skills",
                "skills profile",
                "key skills",
                "technical skills",
                "core skills",
                "key competencies",
                "core competencies",
                "skills & competencies",
            )
            if key in sections
        ),
        "",
    )
    skill_lines = _clean_cv_lines(skills_text)
    trimmed_skill_lines: list[str] = []
    for index, line in enumerate(skill_lines):
        if index + 1 < len(skill_lines) and _is_job_date_line(skill_lines[index + 1]):
            break
        if line.casefold().rstrip(":") in _STOP_HEADINGS:
            break
        trimmed_skill_lines.append(line)
    skills = _dedupe(
        [
            part
            for line in trimmed_skill_lines
            for part in re.split(r"[,|;•·]+", line)
            if 1 < len(part.strip()) < 100
        ],
        limit=30,
    )

    qualification_text = next(
        (
            sections[key]
            for key in (
                "qualifications",
                "education",
                "education/training",
                "education & training",
                "education & qualifications",
                "qualifications & education",
                "certifications",
                "professional qualifications",
            )
            if key in sections
        ),
        "",
    )
    qualification_lines = _clean_cv_lines(qualification_text)
    qualifications: list[CandidateQualification] = []
    pending_institution = ""
    pending_date = ""
    for line in qualification_lines:
        lowered = line.casefold().rstrip(":")
        if lowered in {"interests", "references"}:
            break
        if re.match(rf"(?i)^\s*{_MONTH}\s+\d{{4}}\s*[–—-]", line):
            pending_date = line
            if ":" in line:
                date_part, institution = line.split(":", 1)
                pending_date = date_part.strip()
                pending_institution = institution.strip()
            continue
        if len(line) >= 3:
            qualifications.append(
                CandidateQualification(
                    name=line,
                    institution=pending_institution,
                    date=pending_date,
                )
            )
            pending_institution = ""
            pending_date = ""
        if len(qualifications) >= 12:
            break

    experience_text = next(
        (
            sections[key]
            for key in (
                "experience",
                "employment",
                "employment history",
                "work history",
                "career history",
                "work experience",
                "professional experience",
                "career experience",
            )
            if key in sections
        ),
        "",
    )
    experience_source = _clean_cv_lines(experience_text) if experience_text else lines
    experience = _experience_from_role_date_pairs(experience_source, skills)
    if not experience and experience_text:
        experience_lines = _dedupe(_clean_cv_lines(experience_text), limit=30)
        if experience_lines:
            experience.append(
                CandidateExperience(
                    role="CV experience",
                    highlights=experience_lines[:20],
                    skills=skills[:12],
                )
            )

    summary_text = next(
        (
            sections[key]
            for key in (
                "profile",
                "personal profile",
                "professional profile",
                "objective",
                "career objective",
                "summary",
                "professional summary",
                "about me",
            )
            if key in sections
        ),
        "",
    )
    summary = " ".join(_clean_cv_lines(summary_text))[:1200]
    if not summary:
        safe_lines = [
            line
            for line in lines
            if not re.search(r"\b(?:tel|phone|mobile|email|e-mail)\b", line, re.IGNORECASE)
            and "@" not in line
        ]
        summary = " ".join(safe_lines[:8])[:700]

    return CandidateProfileData(
        summary=summary,
        skills=skills,
        experience=experience,
        qualifications=qualifications,
    )


def _openai_profile(text: str) -> CandidateProfileData | None:
    if not settings.OPENAI_API_KEY:
        return None
    text = "\n".join(_clean_cv_lines(text))
    try:
        from openai import OpenAI

        client = OpenAI(api_key=settings.OPENAI_API_KEY)
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            response_format={"type": "json_object"},
            temperature=0,
            max_tokens=1800,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Extract only facts explicitly present in the candidate CV. "
                        "Do not infer seniority, qualifications, skills, dates, employers or achievements that are not stated. "
                        "Return JSON with keys summary (string), skills (string array), experience "
                        "(array of objects with role, organisation, dates, highlights string array, skills string array), "
                        "and qualifications (array of objects with name, institution, date)."
                    ),
                },
                {"role": "user", "content": text},
            ],
        )
        raw = response.choices[0].message.content or "{}"
        parsed = json.loads(raw)
        return CandidateProfileData.model_validate(parsed)
    except Exception:
        return None


def _has_structured_profile(profile: CandidateProfileData) -> bool:
    return bool(profile.skills or profile.experience or profile.qualifications)


def extract_candidate_profile(text: str) -> tuple[CandidateProfileData, str]:
    ai_profile = _openai_profile(text)
    if ai_profile is not None:
        ai_profile.skills = _dedupe(ai_profile.skills, limit=40)
        for item in ai_profile.experience:
            item.highlights = _dedupe(item.highlights, limit=20)
            item.skills = _dedupe(item.skills, limit=20)
        if _has_structured_profile(ai_profile):
            return ai_profile, "openai"

    fallback = _fallback_profile(text)
    return fallback, "fallback"


def _profile_row_to_response(row: dict[str, Any]) -> CandidateProfileResponse:
    metadata: dict[str, Any] = {}
    try:
        metadata = json.loads(str(row.get("reflection") or "{}"))
    except Exception:
        metadata = {}
    experience = metadata.get("experience") if isinstance(metadata.get("experience"), list) else []
    qualifications = metadata.get("qualifications") if isinstance(metadata.get("qualifications"), list) else []
    return CandidateProfileResponse(
        user_id=str(row.get("user_id") or ""),
        source_filename=str(metadata.get("source_filename") or row.get("task") or ""),
        summary=str(row.get("situation") or ""),
        skills=[str(item) for item in (row.get("skills") or []) if item],
        experience=experience,
        qualifications=qualifications,
        extraction_provider=str(metadata.get("extraction_provider") or "fallback"),
        updated_at=row.get("updated_at") or row.get("created_at"),
    )


def _profile_payload(
    *,
    user_id: str,
    filename: str,
    profile: CandidateProfileData,
    provider: str,
) -> dict[str, Any]:
    metadata = {
        "source_filename": filename[:255],
        "extraction_provider": provider,
        "experience": [item.model_dump() for item in profile.experience],
        "qualifications": [item.model_dump() for item in profile.qualifications],
    }
    return {
        "user_id": user_id,
        "title": "CV profile",
        "situation": profile.summary,
        "task": filename[:255],
        "actions": [],
        "outcome": "",
        "reflection": json.dumps(metadata, ensure_ascii=False),
        "tags": [item.name for item in profile.qualifications][:30],
        "behaviours": [],
        "skills": profile.skills,
        "authority_context": None,
        "source": "cv_profile",
        "confidence": 50,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


@router.get("", response_model=CandidateProfileResponse | None)
async def get_candidate_profile(
    authorization: str | None = Header(None),
) -> CandidateProfileResponse | None:
    user = await verify_supabase_user(authorization)
    result = (
        _db()
        .table("evidence_cards")
        .select("*")
        .eq("user_id", user["id"])
        .eq("source", "cv_profile")
        .execute()
    )
    rows = getattr(result, "data", None) or []
    if not rows:
        return None
    rows = sorted(rows, key=lambda row: str(row.get("updated_at") or ""), reverse=True)
    return _profile_row_to_response(rows[0])


@router.post("/cv-upload", response_model=CandidateProfileResponse)
async def upload_candidate_cv(
    file: UploadFile = File(...),
    authorization: str | None = Header(None),
) -> CandidateProfileResponse:
    user = await verify_supabase_user(authorization)
    filename = Path(file.filename or "cv").name
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise HTTPException(status_code=415, detail="Upload a PDF, DOCX or TXT CV.")

    data = await file.read(MAX_CV_BYTES + 1)
    if len(data) > MAX_CV_BYTES:
        raise HTTPException(status_code=413, detail="CV files must be 5 MB or smaller.")

    text = _extract_text(filename, data)
    profile, provider = extract_candidate_profile(text)
    if not _has_structured_profile(profile):
        raise HTTPException(
            status_code=422,
            detail=(
                "We could read this CV, but could not reliably identify skills, employment history "
                "or qualifications. Try exporting it as a simpler DOCX/PDF or TXT file."
            ),
        )
    payload = _profile_payload(
        user_id=user["id"],
        filename=filename,
        profile=profile,
        provider=provider,
    )

    current = (
        _db()
        .table("evidence_cards")
        .select("id")
        .eq("user_id", user["id"])
        .eq("source", "cv_profile")
        .execute()
    )
    rows = getattr(current, "data", None) or []
    if rows:
        result = (
            _db()
            .table("evidence_cards")
            .update({key: value for key, value in payload.items() if key != "user_id"})
            .eq("id", rows[0]["id"])
            .eq("user_id", user["id"])
            .execute()
        )
    else:
        result = _db().table("evidence_cards").insert(payload).execute()

    saved = getattr(result, "data", None) or []
    if not saved:
        raise HTTPException(status_code=503, detail="Could not save the extracted CV profile.")
    return _profile_row_to_response(saved[0])


@router.delete("")
async def delete_candidate_profile(
    authorization: str | None = Header(None),
) -> dict[str, bool]:
    user = await verify_supabase_user(authorization)
    (
        _db()
        .table("evidence_cards")
        .delete()
        .eq("user_id", user["id"])
        .eq("source", "cv_profile")
        .execute()
    )
    return {"ok": True}
