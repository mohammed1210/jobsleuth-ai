from __future__ import annotations

from datetime import datetime, timezone

from fastapi.testclient import TestClient

from backend.main import app
from routes import candidate_profile
from routes.candidate_profile import CandidateExperience, CandidateProfileData
from routes.vacancy_analysis import _profile_support

client = TestClient(app)


class FakeResult:
    def __init__(self, data):
        self.data = data


class FakeTable:
    def __init__(self):
        self.rows = []
        self.operation = "select"
        self.payload = None
        self.filters = []

    def select(self, *args, **kwargs):
        self.operation = "select"
        self.filters = []
        return self

    def insert(self, payload):
        self.operation = "insert"
        self.payload = payload
        self.filters = []
        return self

    def update(self, payload):
        self.operation = "update"
        self.payload = payload
        self.filters = []
        return self

    def delete(self):
        self.operation = "delete"
        self.filters = []
        return self

    def eq(self, field, value):
        self.filters.append((field, value))
        return self

    def _matches(self, row):
        return all(row.get(field) == value for field, value in self.filters)

    def execute(self):
        now = datetime.now(timezone.utc).isoformat()
        if self.operation == "insert":
            row = {"id": "cv-1", "created_at": now, "updated_at": now, **self.payload}
            self.rows.append(row)
            return FakeResult([dict(row)])
        if self.operation == "update":
            updated = []
            for row in self.rows:
                if self._matches(row):
                    row.update(self.payload)
                    updated.append(dict(row))
            return FakeResult(updated)
        if self.operation == "delete":
            removed = [dict(row) for row in self.rows if self._matches(row)]
            self.rows = [row for row in self.rows if not self._matches(row)]
            return FakeResult(removed)
        return FakeResult([dict(row) for row in self.rows if self._matches(row)])


class FakeClient:
    def __init__(self):
        self.evidence = FakeTable()

    def table(self, name):
        assert name == "evidence_cards"
        return self.evidence


def test_fallback_extracts_cv_sections(monkeypatch):
    monkeypatch.setattr(candidate_profile, "_openai_profile", lambda text: None)
    text = """
PROFILE
Operations professional with border, freight and stakeholder experience.

SKILLS
Stakeholder management
Risk assessment
Data analysis

EXPERIENCE
Border Force Officer
Reviewed freight movements and escalated operational risks.
Worked with internal and external stakeholders.

QUALIFICATIONS
Level 3 Certificate in Investigations
BSc Business Management
"""
    profile, provider = candidate_profile.extract_candidate_profile(text)
    assert provider == "fallback"
    assert "Stakeholder management" in profile.skills
    assert profile.experience
    assert any("Level 3 Certificate" in item.name for item in profile.qualifications)


def test_txt_cv_upload_persists_only_structured_profile(monkeypatch):
    fake = FakeClient()
    monkeypatch.setattr(candidate_profile, "get_supabase_client", lambda: fake)
    monkeypatch.setattr(candidate_profile, "_openai_profile", lambda text: None)
    headers = {"Authorization": "Bearer valid_token"}
    cv = b"""PROFILE
Operations professional with substantial stakeholder experience.

SKILLS
Stakeholder management
Risk assessment

EXPERIENCE
Border Force Officer
Reviewed operational risks and worked with partner agencies.

QUALIFICATIONS
Level 3 Certificate in Investigations
"""

    response = client.post(
        "/candidate-profile/cv-upload",
        headers=headers,
        files={"file": ("candidate.txt", cv, "text/plain")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["source_filename"] == "candidate.txt"
    assert "Stakeholder management" in body["skills"]
    assert fake.evidence.rows[0]["source"] == "cv_profile"
    assert fake.evidence.rows[0]["title"] == "CV profile"
    assert "PROFILE\n" not in fake.evidence.rows[0]["reflection"]

    loaded = client.get("/candidate-profile", headers=headers)
    assert loaded.status_code == 200
    assert loaded.json()["skills"] == body["skills"]


def test_cv_upload_rejects_unsupported_file_type(monkeypatch):
    fake = FakeClient()
    monkeypatch.setattr(candidate_profile, "get_supabase_client", lambda: fake)
    response = client.post(
        "/candidate-profile/cv-upload",
        headers={"Authorization": "Bearer valid_token"},
        files={"file": ("candidate.rtf", b"some cv content" * 20, "application/rtf")},
    )
    assert response.status_code == 415


def test_profile_support_is_capped_at_partial():
    profile = CandidateProfileData(
        summary="Experienced operational professional.",
        skills=["stakeholder management", "risk assessment"],
        experience=[
            CandidateExperience(
                role="Operations Officer",
                organisation="Example",
                highlights=[
                    "Worked with internal and external stakeholders to resolve operational risks.",
                    "Completed risk assessments and documented findings.",
                ],
                skills=["stakeholder management", "risk assessment"],
            )
        ],
        qualifications=[],
    )
    support = _profile_support(
        "Strong stakeholder management and risk assessment experience",
        profile,
    )
    assert support is not None
    assert support["strength"] in {"partial", "weak"}
    assert support["strength"] != "strong"


def test_profile_does_not_fake_people_management_scope():
    profile = CandidateProfileData(
        skills=["security operations", "risk management"],
        experience=[
            CandidateExperience(
                role="Security Officer",
                highlights=["Managed security risks across operational activity."],
                skills=["security operations", "risk management"],
            )
        ],
        qualifications=[],
    )
    support = _profile_support(
        "Experience at management level within security operations",
        profile,
    )
    assert support is None or support["strength"] != "partial"


def test_cv_profile_can_support_qualification_without_becoming_strong():
    profile = CandidateProfileData(
        summary="Operations professional.",
        skills=[],
        experience=[],
        qualifications=[
            {"name": "BSc Business Management", "institution": "Example University", "date": "2020"}
        ],
    )
    support = _profile_support("BSc Business Management qualification", profile)
    assert support is not None
    assert support["strength"] == "partial"


def test_cv_signal_cannot_clear_explicit_blocker():
    response = client.post(
        "/vacancy-analysis",
        headers={"Authorization": "Bearer valid_token"},
        json={
            "job": {"title": "Security Manager"},
            "requirements": [
                {
                    "text": "Strong stakeholder management and risk assessment experience",
                    "category": "essential",
                    "blocker": True,
                }
            ],
            "evidence_cards": [],
            "candidate_profile": {
                "summary": "Operational professional.",
                "skills": ["stakeholder management", "risk assessment"],
                "experience": [
                    {
                        "role": "Operations Officer",
                        "organisation": "Example",
                        "dates": "",
                        "highlights": [
                            "Worked with stakeholders and completed risk assessments."
                        ],
                        "skills": ["stakeholder management", "risk assessment"],
                    }
                ],
                "qualifications": [],
            },
            "practical_issues": [],
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["decision"] == "SKIP"
    assert body["requirements"][0]["match_strength"] in {"weak", "missing"}
    assert body["requirements"][0]["profile_support"] is not None


def test_docx_extraction_includes_table_cells():
    import io

    from docx import Document

    document = Document()
    document.add_paragraph("Candidate profile")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "SKILLS"
    table.cell(1, 0).text = "Stakeholder management\nRisk assessment"
    buffer = io.BytesIO()
    document.save(buffer)

    text = candidate_profile._extract_docx(buffer.getvalue())
    assert "Stakeholder management" in text
    assert "Risk assessment" in text


def test_fallback_recognises_common_cv_heading_variants(monkeypatch):
    monkeypatch.setattr(candidate_profile, "_openai_profile", lambda text: None)
    text = """
PROFESSIONAL PROFILE
Operational professional with public-facing enforcement experience.

KEY COMPETENCIES
Stakeholder Management
Report Writing
Risk Assessment

PROFESSIONAL EXPERIENCE
Border Force Officer
Gathered intelligence, assessed risk and worked with partner agencies.

EDUCATION & QUALIFICATIONS
Level 3 Certificate in Investigations
"""
    profile, provider = candidate_profile.extract_candidate_profile(text)
    assert provider == "fallback"
    assert "Stakeholder Management" in profile.skills
    assert profile.experience
    assert any("Level 3 Certificate" in item.name for item in profile.qualifications)


def test_upload_rejects_empty_structured_profile(monkeypatch):
    fake = FakeClient()
    monkeypatch.setattr(candidate_profile, "get_supabase_client", lambda: fake)
    monkeypatch.setattr(
        candidate_profile,
        "extract_candidate_profile",
        lambda _text: (CandidateProfileData(summary="Readable CV text only."), "fallback"),
    )
    response = client.post(
        "/candidate-profile/cv-upload",
        headers={"Authorization": "Bearer valid_token"},
        files={"file": ("candidate.txt", b"A readable CV with enough text to pass the file extraction threshold. " * 4, "text/plain")},
    )
    assert response.status_code == 422
    assert "could not reliably identify" in response.json()["detail"].lower()
    assert fake.evidence.rows == []
